#!/usr/bin/env python3
"""
엑셀(.xlsx) 밸런스 시트 <-> DungeonVM 밸런스 JSON(BalanceData) 변환기.

기획자는 balance.xlsx를 열어 각 시트의 값만 수정하면 되고, 이 스크립트가
DungeonVM.Core / DungeonVM.Simulator가 그대로 읽는 JSON(스키마: DungeonVM.Core/Balance/BalanceData.cs)으로
변환해준다.

사용법:
    # 1) 최초 1회: 현재 DefaultBalance.json 값으로 채워진 템플릿 엑셀 생성
    python xlsx_to_balance_json.py template balance.xlsx

    # 2) 기획자가 balance.xlsx를 수정한 뒤, JSON으로 변환
    #    --out을 생략하면 DungeonVM.Core/Balance/DefaultBalance.json(기본값)에 바로 반영된다.
    python xlsx_to_balance_json.py convert balance.xlsx
    python xlsx_to_balance_json.py convert balance.xlsx --out my_override.json

    # 3) 변환된 JSON으로 즉시 시뮬레이터 재실행(선택)
    python xlsx_to_balance_json.py convert balance.xlsx --simulate 500
"""

from __future__ import annotations

import argparse
import json
import subprocess
import sys
from pathlib import Path

if sys.platform == "win32":
    # Windows 콘솔 코드페이지(cp949 등)와 무관하게 한글 출력이 항상 UTF-8로 나가도록 강제한다.
    for _stream in (sys.stdout, sys.stderr):
        try:
            _stream.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from balance_schema import (
    ARMOR_MARKET_VALUES_COLUMNS,
    ARMOR_MARKET_VALUES_SHEET,
    ARMOR_ROLL_RANGES_COLUMNS,
    ARMOR_ROLL_RANGES_SHEET,
    ARMOR_RARITIES,
    CHARACTER_SLOT_UNLOCK_COSTS_COLUMNS,
    CHARACTER_SLOT_UNLOCK_COSTS_SHEET,
    MERGE_GRID_UNLOCK_COSTS_COLUMNS,
    MERGE_GRID_UNLOCK_COSTS_SHEET,
    SCALAR_FIELDS,
    SCALAR_FIELDS_BY_KEY,
    SCALARS_COLUMNS,
    SCALARS_SHEET,
    VM_UPGRADE_COSTS_COLUMNS,
    VM_UPGRADE_COSTS_SHEET,
    WEAPON_LEVEL_MULTIPLIERS_COLUMNS,
    WEAPON_LEVEL_MULTIPLIERS_SHEET,
    WEAPON_TYPES,
    WEAPONS_COLUMNS,
    WEAPONS_SHEET,
)

# backend/tools/balance_pipeline/이 파일 위치이므로 parents[2] = backend/ (Unity 프로젝트가 있는 저장소 루트가 아님).
BACKEND_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_BALANCE_JSON = BACKEND_ROOT / "DungeonVM.Core" / "Balance" / "DefaultBalance.json"
SIMULATOR_PROJECT = BACKEND_ROOT / "DungeonVM.Simulator" / "DungeonVM.Simulator.csproj"


# ---------------------------------------------------------------------------
# 템플릿 생성: 현재 DefaultBalance.json 값을 그대로 엑셀 시트에 채워 넣는다.
# ---------------------------------------------------------------------------

def make_template(out_path: Path, seed_json_path: Path = DEFAULT_BALANCE_JSON) -> None:
    seed = json.loads(seed_json_path.read_text(encoding="utf-8")) if seed_json_path.exists() else {}

    wb = Workbook()
    wb.remove(wb.active)

    _write_scalars_sheet(wb, seed)
    _write_weapons_sheet(wb, seed)
    _write_weapon_level_multipliers_sheet(wb, seed)
    _write_armor_roll_ranges_sheet(wb, seed)
    _write_armor_market_values_sheet(wb, seed)
    _write_vm_upgrade_costs_sheet(wb, seed)
    _write_merge_grid_unlock_costs_sheet(wb, seed)
    _write_character_slot_unlock_costs_sheet(wb, seed)

    wb.save(out_path)
    print(f"템플릿 생성 완료: {out_path}")


def _write_scalars_sheet(wb: Workbook, seed: dict) -> None:
    ws = wb.create_sheet(SCALARS_SHEET)
    ws.append(SCALARS_COLUMNS)
    for f in SCALAR_FIELDS:
        section = seed.get(_camel(f.section), {})
        value = section.get(_camel(f.field))
        ws.append([f.key, value, f.description_kr])
    _autosize(ws)


def _write_weapons_sheet(wb: Workbook, seed: dict) -> None:
    ws = wb.create_sheet(WEAPONS_SHEET)
    ws.append(WEAPONS_COLUMNS)
    table = seed.get("weapons", {}).get("table", {})
    for weapon_type in WEAPON_TYPES:
        entry = table.get(weapon_type, {})
        ws.append([
            weapon_type,
            entry.get("row", "Front"),
            entry.get("baseDamage"),
            entry.get("attacksPerSecond"),
            entry.get("bonusHealth"),
            entry.get("pullAggro"),
        ])
    _autosize(ws)


def _write_armor_roll_ranges_sheet(wb: Workbook, seed: dict) -> None:
    ws = wb.create_sheet(ARMOR_ROLL_RANGES_SHEET)
    ws.append(ARMOR_ROLL_RANGES_COLUMNS)
    ranges = seed.get("armor", {}).get("rollRanges", {})
    for rarity in ARMOR_RARITIES:
        r = ranges.get(rarity, {})
        ws.append([
            rarity, r.get("minHp"), r.get("maxHp"),
            r.get("minAtk"), r.get("maxAtk"),
            r.get("minDodge"), r.get("maxDodge"),
        ])
    _autosize(ws)


def _write_armor_market_values_sheet(wb: Workbook, seed: dict) -> None:
    ws = wb.create_sheet(ARMOR_MARKET_VALUES_SHEET)
    ws.append(ARMOR_MARKET_VALUES_COLUMNS)
    values = seed.get("currency", {}).get("armorMarketValues", {})
    for rarity in ARMOR_RARITIES:
        ws.append([rarity, values.get(rarity)])
    _autosize(ws)


def _write_vm_upgrade_costs_sheet(wb: Workbook, seed: dict) -> None:
    ws = wb.create_sheet(VM_UPGRADE_COSTS_SHEET)
    ws.append(VM_UPGRADE_COSTS_COLUMNS)
    costs = seed.get("vendingMachine", {}).get("upgradeGoldCosts", [])
    for level, cost in enumerate(costs, start=1):
        ws.append([level, cost])
    _autosize(ws)


def _write_merge_grid_unlock_costs_sheet(wb: Workbook, seed: dict) -> None:
    ws = wb.create_sheet(MERGE_GRID_UNLOCK_COSTS_SHEET)
    ws.append(MERGE_GRID_UNLOCK_COSTS_COLUMNS)
    costs = seed.get("mergeGrid", {}).get("unlockGoldCosts", [])
    for block, cost in enumerate(costs, start=1):
        ws.append([block, cost])
    _autosize(ws)


def _write_weapon_level_multipliers_sheet(wb: Workbook, seed: dict) -> None:
    ws = wb.create_sheet(WEAPON_LEVEL_MULTIPLIERS_SHEET)
    ws.append(WEAPON_LEVEL_MULTIPLIERS_COLUMNS)
    multipliers = seed.get("weapons", {}).get("levelMultipliers", [])
    for level, multiplier in enumerate(multipliers, start=1):
        ws.append([level, multiplier])
    _autosize(ws)


def _write_character_slot_unlock_costs_sheet(wb: Workbook, seed: dict) -> None:
    ws = wb.create_sheet(CHARACTER_SLOT_UNLOCK_COSTS_SHEET)
    ws.append(CHARACTER_SLOT_UNLOCK_COSTS_COLUMNS)
    starting = seed.get("characterSlots", {}).get("startingSlots", 2)
    costs = seed.get("characterSlots", {}).get("unlockGoldCosts", [])
    for offset, cost in enumerate(costs, start=1):
        ws.append([starting + offset, cost])
    _autosize(ws)


def _autosize(ws: Worksheet, min_width: int = 10) -> None:
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=0)
        ws.column_dimensions[col_cells[0].column_letter].width = max(min_width, length + 2)


# ---------------------------------------------------------------------------
# 변환: 엑셀 -> BalanceData JSON
# ---------------------------------------------------------------------------

class ConversionError(Exception):
    pass


def convert(xlsx_path: Path) -> dict:
    wb = load_workbook(xlsx_path, data_only=True)
    errors: list[str] = []

    data: dict = {}
    _read_scalars_sheet(wb, data, errors)
    _read_weapons_sheet(wb, data, errors)
    _read_weapon_level_multipliers_sheet(wb, data, errors)
    _read_armor_roll_ranges_sheet(wb, data, errors)
    _read_armor_market_values_sheet(wb, data, errors)
    _read_vm_upgrade_costs_sheet(wb, data, errors)
    _read_merge_grid_unlock_costs_sheet(wb, data, errors)
    _read_character_slot_unlock_costs_sheet(wb, data, errors)

    if errors:
        raise ConversionError("\n".join(errors))

    return data


def _rows(ws: Worksheet):
    header = [str(c.value).strip() if c.value is not None else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
    for row in ws.iter_rows(min_row=2):
        values = [c.value for c in row]
        if all(v is None for v in values):
            continue
        yield dict(zip(header, values))


def _sheet(wb, name: str, errors: list[str]) -> Worksheet | None:
    if name not in wb.sheetnames:
        errors.append(f"[{name}] 시트가 없습니다.")
        return None
    return wb[name]


def _set_section_field(data: dict, section: str, field: str, value) -> None:
    data.setdefault(_camel(section), {})[_camel(field)] = value


def _camel(pascal: str) -> str:
    return pascal[0].lower() + pascal[1:] if pascal else pascal


def _read_scalars_sheet(wb, data: dict, errors: list[str]) -> None:
    ws = _sheet(wb, SCALARS_SHEET, errors)
    if ws is None:
        return

    seen: set[str] = set()
    for row in _rows(ws):
        key = row.get("Key")
        if key is None or str(key).strip() == "":
            continue
        key = str(key).strip()
        value = row.get("Value")

        field = SCALAR_FIELDS_BY_KEY.get(key)
        if field is None:
            errors.append(f"[{SCALARS_SHEET}] 알 수 없는 Key '{key}' (오타 여부를 확인하세요)")
            continue
        if value is None:
            errors.append(f"[{SCALARS_SHEET}] '{key}' 값이 비어 있습니다.")
            continue

        try:
            typed_value = field.value_type(value)
        except (TypeError, ValueError):
            errors.append(f"[{SCALARS_SHEET}] '{key}' 값 '{value}'을(를) {field.value_type.__name__}(으)로 변환할 수 없습니다.")
            continue

        _set_section_field(data, field.section, field.field, typed_value)
        seen.add(key)

    missing = [f.key for f in SCALAR_FIELDS if f.key not in seen]
    if missing:
        errors.append(f"[{SCALARS_SHEET}] 다음 Key가 누락되었습니다: {', '.join(missing)}")


def _read_weapons_sheet(wb, data: dict, errors: list[str]) -> None:
    ws = _sheet(wb, WEAPONS_SHEET, errors)
    if ws is None:
        return

    table: dict = {}
    for row in _rows(ws):
        weapon_type = row.get("Type")
        if weapon_type is None:
            continue
        weapon_type = str(weapon_type).strip()
        if weapon_type not in WEAPON_TYPES:
            errors.append(f"[{WEAPONS_SHEET}] 알 수 없는 Type '{weapon_type}' (허용값: {', '.join(WEAPON_TYPES)})")
            continue
        try:
            table[weapon_type] = {
                "row": str(row["Row"]).strip(),
                "baseDamage": float(row["BaseDamage"]),
                "attacksPerSecond": float(row["AttacksPerSecond"]),
                "bonusHealth": float(row["BonusHealth"]),
                "pullAggro": float(row["PullAggro"]),
            }
        except (KeyError, TypeError, ValueError) as e:
            errors.append(f"[{WEAPONS_SHEET}] '{weapon_type}' 행 값이 잘못되었습니다: {e}")

    missing = [t for t in WEAPON_TYPES if t not in table]
    if missing:
        errors.append(f"[{WEAPONS_SHEET}] 다음 무기 타입이 누락되었습니다: {', '.join(missing)}")

    data.setdefault("weapons", {})["table"] = table


def _read_armor_roll_ranges_sheet(wb, data: dict, errors: list[str]) -> None:
    ws = _sheet(wb, ARMOR_ROLL_RANGES_SHEET, errors)
    if ws is None:
        return

    ranges: dict = {}
    for row in _rows(ws):
        rarity = row.get("Rarity")
        if rarity is None:
            continue
        rarity = str(rarity).strip()
        if rarity not in ARMOR_RARITIES:
            errors.append(f"[{ARMOR_ROLL_RANGES_SHEET}] 알 수 없는 Rarity '{rarity}' (허용값: {', '.join(ARMOR_RARITIES)})")
            continue
        try:
            ranges[rarity] = {
                "minHp": float(row["MinHp"]), "maxHp": float(row["MaxHp"]),
                "minAtk": float(row["MinAtk"]), "maxAtk": float(row["MaxAtk"]),
                "minDodge": float(row["MinDodge"]), "maxDodge": float(row["MaxDodge"]),
            }
        except (KeyError, TypeError, ValueError) as e:
            errors.append(f"[{ARMOR_ROLL_RANGES_SHEET}] '{rarity}' 행 값이 잘못되었습니다: {e}")

    missing = [r for r in ARMOR_RARITIES if r not in ranges]
    if missing:
        errors.append(f"[{ARMOR_ROLL_RANGES_SHEET}] 다음 등급이 누락되었습니다: {', '.join(missing)}")

    data.setdefault("armor", {})["rollRanges"] = ranges


def _read_armor_market_values_sheet(wb, data: dict, errors: list[str]) -> None:
    ws = _sheet(wb, ARMOR_MARKET_VALUES_SHEET, errors)
    if ws is None:
        return

    values: dict = {}
    for row in _rows(ws):
        rarity = row.get("Rarity")
        if rarity is None:
            continue
        rarity = str(rarity).strip()
        if rarity not in ARMOR_RARITIES:
            errors.append(f"[{ARMOR_MARKET_VALUES_SHEET}] 알 수 없는 Rarity '{rarity}'")
            continue
        try:
            values[rarity] = int(row["Value"])
        except (KeyError, TypeError, ValueError) as e:
            errors.append(f"[{ARMOR_MARKET_VALUES_SHEET}] '{rarity}' 값이 잘못되었습니다: {e}")

    data.setdefault("currency", {})["armorMarketValues"] = values


def _read_vm_upgrade_costs_sheet(wb, data: dict, errors: list[str]) -> None:
    ws = _sheet(wb, VM_UPGRADE_COSTS_SHEET, errors)
    if ws is None:
        return

    entries = []
    for row in _rows(ws):
        if row.get("Level") is None:
            continue
        try:
            entries.append((int(row["Level"]), int(row["GoldCost"])))
        except (KeyError, TypeError, ValueError) as e:
            errors.append(f"[{VM_UPGRADE_COSTS_SHEET}] 행 값이 잘못되었습니다: {e}")

    entries.sort(key=lambda t: t[0])
    data.setdefault("vendingMachine", {})["upgradeGoldCosts"] = [cost for _, cost in entries]


def _read_merge_grid_unlock_costs_sheet(wb, data: dict, errors: list[str]) -> None:
    ws = _sheet(wb, MERGE_GRID_UNLOCK_COSTS_SHEET, errors)
    if ws is None:
        return

    entries = []
    for row in _rows(ws):
        if row.get("Block") is None:
            continue
        try:
            entries.append((int(row["Block"]), int(row["GoldCost"])))
        except (KeyError, TypeError, ValueError) as e:
            errors.append(f"[{MERGE_GRID_UNLOCK_COSTS_SHEET}] 행 값이 잘못되었습니다: {e}")

    entries.sort(key=lambda t: t[0])
    data.setdefault("mergeGrid", {})["unlockGoldCosts"] = [cost for _, cost in entries]


def _read_weapon_level_multipliers_sheet(wb, data: dict, errors: list[str]) -> None:
    ws = _sheet(wb, WEAPON_LEVEL_MULTIPLIERS_SHEET, errors)
    if ws is None:
        return

    entries = []
    for row in _rows(ws):
        if row.get("Level") is None:
            continue
        try:
            entries.append((int(row["Level"]), float(row["Multiplier"])))
        except (KeyError, TypeError, ValueError) as e:
            errors.append(f"[{WEAPON_LEVEL_MULTIPLIERS_SHEET}] 행 값이 잘못되었습니다: {e}")

    entries.sort(key=lambda t: t[0])
    data.setdefault("weapons", {})["levelMultipliers"] = [multiplier for _, multiplier in entries]


def _read_character_slot_unlock_costs_sheet(wb, data: dict, errors: list[str]) -> None:
    ws = _sheet(wb, CHARACTER_SLOT_UNLOCK_COSTS_SHEET, errors)
    if ws is None:
        return

    entries = []
    for row in _rows(ws):
        if row.get("Slot") is None:
            continue
        try:
            entries.append((int(row["Slot"]), int(row["GoldCost"])))
        except (KeyError, TypeError, ValueError) as e:
            errors.append(f"[{CHARACTER_SLOT_UNLOCK_COSTS_SHEET}] 행 값이 잘못되었습니다: {e}")

    entries.sort(key=lambda t: t[0])
    data.setdefault("characterSlots", {})["unlockGoldCosts"] = [cost for _, cost in entries]


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="command", required=True)

    p_template = sub.add_parser("template", help="현재 DefaultBalance.json 값으로 채워진 엑셀 템플릿 생성")
    p_template.add_argument("out_xlsx", type=Path, help="생성할 .xlsx 경로")
    p_template.add_argument("--seed", type=Path, default=DEFAULT_BALANCE_JSON, help="템플릿 초기값으로 쓸 JSON (기본: DefaultBalance.json)")

    p_convert = sub.add_parser("convert", help="엑셀을 밸런스 JSON으로 변환")
    p_convert.add_argument("xlsx", type=Path, help="입력 .xlsx 경로")
    p_convert.add_argument("--out", type=Path, default=DEFAULT_BALANCE_JSON, help="출력 JSON 경로 (기본: DungeonVM.Core/Balance/DefaultBalance.json)")
    p_convert.add_argument("--simulate", type=int, metavar="RUNS_PER_BOT", help="변환 직후 이 JSON으로 시뮬레이터를 RUNS_PER_BOT회씩 실행")

    args = parser.parse_args(argv)

    if args.command == "template":
        make_template(args.out_xlsx, args.seed)
        return 0

    if args.command == "convert":
        try:
            data = convert(args.xlsx)
        except ConversionError as e:
            print("변환 실패 - 엑셀 값을 확인하세요:\n", file=sys.stderr)
            print(str(e), file=sys.stderr)
            return 1

        args.out.parent.mkdir(parents=True, exist_ok=True)
        args.out.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"변환 완료: {args.xlsx} -> {args.out}")

        if args.simulate is not None:
            _run_simulator(args.out.resolve(), args.simulate)

        return 0

    return 1


def _run_simulator(balance_json: Path, runs_per_bot: int) -> None:
    print(f"\n시뮬레이터 실행 중... ({runs_per_bot}회/봇, 밸런스={balance_json})")
    cmd = [
        "dotnet", "run", "--project", str(SIMULATOR_PROJECT), "--",
        str(runs_per_bot), "--balance", str(balance_json),
    ]
    subprocess.run(cmd, check=False, cwd=BACKEND_ROOT)


if __name__ == "__main__":
    raise SystemExit(main())
